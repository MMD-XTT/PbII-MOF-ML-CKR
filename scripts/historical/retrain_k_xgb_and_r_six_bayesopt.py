import json
import shutil
from pathlib import Path
PACKAGE_ROOT = Path(__file__).resolve().parents[2]

import joblib
import numpy as np
import optuna
import pandas as pd
import shap
from catboost import CatBoostRegressor, Pool
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.base import clone
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import ExtraTreesRegressor, RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.linear_model import Ridge
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.model_selection import KFold, cross_val_predict, cross_val_score
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder
from sklearn.svm import SVR
from xgboost import XGBRegressor

optuna.logging.set_verbosity(optuna.logging.WARNING)

DATA_DIR = PACKAGE_ROOT / 'data'
K_ROOT = PACKAGE_ROOT / 'outputs/historical_training/K'
R_ROOT = PACKAGE_ROOT / 'outputs/historical_training/R'
CV = KFold(n_splits=5, shuffle=True, random_state=2026)
N_TRIALS = 30
SEED = 42

BASE_BINARY = [
    "O_donor", "N_donor", "S_donor", "P_donor", "carboxyl", "amino",
    "thiol", "sulfonic", "phosphonic",
]
PORE = ["BET_m2_g", "total_pore_volume_cm3_g", "average_pore_diameter_nm"]


def add_multihot(df, column, prefix):
    tokens = sorted({
        token.strip()
        for value in df[column].fillna("none")
        for token in str(value).split(";")
        if token.strip()
    })
    names = []
    for token in tokens:
        name = f"{prefix}__{token}"
        df[name] = df[column].fillna("none").map(
            lambda value: int(token in {
                x.strip() for x in str(value).split(";") if x.strip()
            })
        )
        names.append(name)
    return names


def prepare(source, target, conditions):
    df = pd.read_excel(source)
    df[target] = pd.to_numeric(df[target], errors="coerce")
    df = df[df[target].notna()].copy().reset_index(drop=True)
    multi = (
        add_multihot(df, "metal_node_elements", "metal")
        + add_multihot(df, "Linker family", "linker")
        + add_multihot(df, "modulator", "modulator")
    )
    numeric = BASE_BINARY + multi + PORE + conditions
    features = ["main"] + numeric
    x = df[features].copy()
    for col in numeric:
        x[col] = pd.to_numeric(x[col], errors="coerce").fillna(0)
    return df, x, df[target].to_numpy(dtype=float), numeric, features


def make_pipeline(numeric, estimator):
    preprocess = ColumnTransformer([
        ("main", Pipeline([
            ("impute", SimpleImputer(strategy="constant", fill_value="none")),
            ("encode", OneHotEncoder(handle_unknown="ignore")),
        ]), ["main"]),
        ("numeric", SimpleImputer(strategy="median", add_indicator=True), numeric),
    ])
    return Pipeline([("preprocess", preprocess), ("model", estimator)])


def suggest_params(trial, name):
    if name == "Ridge":
        return {"alpha": trial.suggest_float("alpha", 0.01, 1000, log=True)}
    if name == "SVR":
        return {
            "C": trial.suggest_float("C", 0.1, 2000, log=True),
            "epsilon": trial.suggest_float("epsilon", 0.001, 10, log=True),
            "gamma": trial.suggest_float("gamma", 1e-4, 10, log=True),
        }
    if name in {"RF", "ExtraTrees"}:
        return {
            "n_estimators": trial.suggest_int("n_estimators", 300, 1000, step=100),
            "max_depth": trial.suggest_int("max_depth", 4, 30),
            "min_samples_leaf": trial.suggest_int("min_samples_leaf", 1, 8),
            "max_features": trial.suggest_float("max_features", 0.3, 1.0),
            "max_samples": trial.suggest_float("max_samples", 0.7, 1.0),
        }
    if name == "XGBoost":
        return {
            "n_estimators": trial.suggest_int("n_estimators", 300, 1000, step=100),
            "max_depth": trial.suggest_int("max_depth", 3, 10),
            "learning_rate": trial.suggest_float("learning_rate", 0.01, 0.2, log=True),
            "subsample": trial.suggest_float("subsample", 0.6, 1.0),
            "colsample_bytree": trial.suggest_float("colsample_bytree", 0.5, 1.0),
            "min_child_weight": trial.suggest_int("min_child_weight", 1, 10),
            "reg_alpha": trial.suggest_float("reg_alpha", 1e-4, 5, log=True),
            "reg_lambda": trial.suggest_float("reg_lambda", 0.1, 20, log=True),
        }
    return {
        "iterations": trial.suggest_int("iterations", 300, 1000, step=100),
        "depth": trial.suggest_int("depth", 4, 10),
        "learning_rate": trial.suggest_float("learning_rate", 0.01, 0.2, log=True),
        "l2_leaf_reg": trial.suggest_float("l2_leaf_reg", 0.1, 20, log=True),
        "random_strength": trial.suggest_float("random_strength", 0, 2),
        "bagging_temperature": trial.suggest_float("bagging_temperature", 0, 5),
    }


def make_estimator(name, params):
    if name == "Ridge":
        return Ridge(**params)
    if name == "SVR":
        return SVR(kernel="rbf", **params)
    if name == "RF":
        return RandomForestRegressor(random_state=SEED, n_jobs=-1, bootstrap=True, **params)
    if name == "ExtraTrees":
        return ExtraTreesRegressor(random_state=SEED, n_jobs=-1, bootstrap=True, **params)
    if name == "XGBoost":
        return XGBRegressor(objective="reg:squarederror", random_state=SEED, n_jobs=-1, **params)
    return CatBoostRegressor(
        loss_function="RMSE", random_seed=SEED, verbose=False, thread_count=-1, **params
    )


def group_for(encoded):
    if encoded.startswith("main__main_"):
        return "M"
    if encoded.startswith("numeric__metal__"):
        return "metal_node_elements"
    if encoded.startswith("numeric__linker__"):
        return "Linker family"
    if encoded.startswith("numeric__modulator__"):
        return "modulator"
    return encoded.replace("numeric__", "")


def component_for(encoded):
    if encoded.startswith("main__main_"):
        return encoded.removeprefix("main__main_")
    for prefix in ("numeric__metal__", "numeric__linker__", "numeric__modulator__"):
        if encoded.startswith(prefix):
            return encoded.removeprefix(prefix)
    return encoded.replace("numeric__", "")


def write_shap_book(summary, group_tables, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Overall SHAP"
    all_sheets = [(ws, summary)] + [(wb.create_sheet(name), frame) for name, frame in group_tables.items()]
    fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet, frame in all_sheets:
        for row in [list(frame.columns)] + frame.values.tolist():
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(name="Times New Roman", bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Times New Roman")
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        for col in range(1, sheet.max_column + 1):
            width = max(len(str(sheet.cell(row=r, column=col).value or "")) for r in range(1, sheet.max_row + 1)) + 2
            sheet.column_dimensions[get_column_letter(col)].width = min(max(width, 14), 34)
    wb.save(path)


def write_tree_shap(final_pipe, x, folder, name):
    transformed = final_pipe.named_steps["preprocess"].transform(x)
    encoded = list(final_pipe.named_steps["preprocess"].get_feature_names_out())
    model = final_pipe.named_steps["model"]
    if name == "XGBoost":
        values = model.get_booster().predict(
            __import__("xgboost").DMatrix(transformed), pred_contribs=True
        )[:, :-1]
    elif name == "CatBoost":
        values = model.get_feature_importance(Pool(transformed), type="ShapValues")[:, :-1]
    else:
        values = shap.TreeExplainer(model).shap_values(transformed)
    values = np.asarray(values)
    table = pd.DataFrame({
        "encoded_feature": encoded,
        "component": [component_for(item) for item in encoded],
        "feature_group": [group_for(item) for item in encoded],
        "mean_abs_shap": np.abs(values).mean(axis=0),
    })
    summary = table.groupby("feature_group", as_index=False)["mean_abs_shap"].sum().sort_values(
        "mean_abs_shap", ascending=False
    ).reset_index(drop=True)
    summary.insert(0, "rank", np.arange(1, len(summary) + 1))
    summary["shap_percentage"] = 100 * summary["mean_abs_shap"] / summary["mean_abs_shap"].sum()
    group_tables = {}
    for group, frame in table.groupby("feature_group"):
        if group in {"M", "metal_node_elements", "Linker family", "modulator"}:
            output = frame[["component", "mean_abs_shap"]].sort_values(
                "mean_abs_shap", ascending=False
            ).reset_index(drop=True)
            output.insert(0, "rank_within_group", np.arange(1, len(output) + 1))
            output["share_within_group_percent"] = 100 * output["mean_abs_shap"] / output["mean_abs_shap"].sum()
            group_tables[group] = output
    summary.to_csv(folder / "SHAP_分组占比.csv", index=False, encoding="utf-8-sig")
    write_shap_book(summary, group_tables, folder / "SHAP_分组占比.xlsx")


def train_one(df, x, y, numeric, features, name, folder, task, target, target_label):
    temporary = folder.with_name(folder.name + "__reoptimizing")
    if temporary.exists():
        shutil.rmtree(temporary)
    temporary.mkdir(parents=True)
    def objective(trial):
        pipeline = make_pipeline(numeric, make_estimator(name, suggest_params(trial, name)))
        return float(cross_val_score(pipeline, x, y, cv=CV, scoring="r2", n_jobs=1).mean())
    study = optuna.create_study(direction="maximize", sampler=optuna.samplers.TPESampler(seed=SEED))
    study.optimize(objective, n_trials=N_TRIALS, show_progress_bar=False)
    best = study.best_params
    pipeline = make_pipeline(numeric, make_estimator(name, best))
    oof = cross_val_predict(pipeline, x, y, cv=CV, method="predict", n_jobs=1)
    folds = []
    for fold, (train_idx, test_idx) in enumerate(CV.split(x), 1):
        model = clone(pipeline).fit(x.iloc[train_idx], y[train_idx])
        prediction = model.predict(x.iloc[test_idx])
        folds.append({
            "fold": fold,
            "r2": float(r2_score(y[test_idx], prediction)),
            "mae": float(mean_absolute_error(y[test_idx], prediction)),
            "rmse": float(mean_squared_error(y[test_idx], prediction) ** 0.5),
            "n_validation": int(len(test_idx)),
        })
    metrics = {
        "model": name,
        "dataset": task,
        "target": target,
        "target_label": target_label,
        "rows": int(len(df)),
        "papers": int(df["paper_id"].nunique()),
        "materials": int(df["material_id"].nunique()),
        "r2_oof": float(r2_score(y, oof)),
        "mae_oof": float(mean_absolute_error(y, oof)),
        "rmse_oof": float(mean_squared_error(y, oof) ** 0.5),
        "r2_fold_mean": float(np.mean([row["r2"] for row in folds])),
        "r2_fold_sd": float(np.std([row["r2"] for row in folds], ddof=1)),
        "best_cv_r2": float(study.best_value),
        "validation": "ordinary random 5-fold KFold, shuffle=True, random_state=2026",
        "optimization": "new Optuna TPE Bayesian optimization after removing T; same 5-fold CV",
        "n_trials": N_TRIALS,
        "numeric_transform": "none; target and conditions retain original units",
        "T_feature": "removed",
    }
    pd.DataFrame({
        "paper_id": df["paper_id"], "material_id": df["material_id"],
        "observed": y, "oof_predicted": oof, "residual": y - oof,
    }).to_csv(temporary / "oof_predictions_5fold.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(folds).to_csv(temporary / "fold_metrics.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(study.trials_dataframe()).to_csv(temporary / "bayesopt_history.csv", index=False, encoding="utf-8-sig")
    (temporary / "metrics.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    (temporary / "parameters.json").write_text(json.dumps({
        "categorical_feature": ["main"], "numeric_features": numeric,
        "all_input_features": features, "target": target,
        "best_model_parameters": best,
        "optimization": metrics["optimization"], "validation": metrics["validation"],
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    pipeline.fit(x, y)
    joblib.dump({"model": pipeline, "feature_columns": features, "scaler": None}, temporary / "final_pipeline.joblib")
    if name in {"RF", "ExtraTrees", "XGBoost", "CatBoost"}:
        write_tree_shap(pipeline, x, temporary, name)
    if folder.exists():
        shutil.rmtree(folder)
    temporary.replace(folder)
    print(f"{task} {name}: OOF R2={metrics['r2_oof']:.4f}, rows={metrics['rows']}", flush=True)
    return metrics


def main():
    # K: rerun solely XGBoost, using exactly all labeled K rows after removing T.
    k_df, k_x, k_y, k_numeric, k_features = prepare(
        DATA_DIR / 'K模型_qt统一.xlsx', "csv_qt",
        ["csv_t", "csv_pH", "csv_temperature_K", "MOF_dose_g_L"],
    )
    k_metrics = train_one(
        k_df, k_x, k_y, k_numeric, k_features, "XGBoost",
        K_ROOT / "XGBoost_qt统一_去T_重跑", "K", "csv_qt", "qt (mg/g)",
    )
    (K_ROOT / "XGBoost_qt统一_去T_重跑_summary.json").write_text(
        json.dumps(k_metrics, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # R: six fresh, independently optimized no-T regressors on the same retained rows/folds.
    r_df, r_x, r_y, r_numeric, r_features = prepare(
        DATA_DIR / 'R_model.xlsx', "csv_retention_percent",
        ["csv_cycle", "csv_pH", "csv_temperature_K", "MOF_dose_g_L"],
    )
    results = []
    for name in ["Ridge", "SVR", "RF", "ExtraTrees", "XGBoost", "CatBoost"]:
        results.append(train_one(
            r_df, r_x, r_y, r_numeric, r_features, name,
            R_ROOT / f"{name}_RetentionPercent_去T_CV5", "R", "csv_retention_percent",
            "retention percent (%)",
        ))
    ranking = pd.DataFrame(sorted(results, key=lambda row: row["r2_oof"], reverse=True))
    ranking.to_csv(R_ROOT / "RetentionPercent_去T_CV5_重新贝叶斯优化_模型对比.csv", index=False, encoding="utf-8-sig")
    (R_ROOT / "RetentionPercent_去T_CV5_重新贝叶斯优化_模型对比.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
