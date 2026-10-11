"""Reusable six-regressor trainer for a no-T MOF feature dataset."""

import argparse
import json
from pathlib import Path

import joblib
import numpy as np
import optuna
import pandas as pd
import shap
from catboost import CatBoostRegressor, Pool
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.base import clone
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import ExtraTreesRegressor, RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.linear_model import Ridge
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.model_selection import KFold, cross_val_predict, cross_val_score
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder
from sklearn.svm import SVR
from xgboost import XGBRegressor

optuna.logging.set_verbosity(optuna.logging.WARNING)
SEED = 42
CV_SEED = 2026
BINARY = ["O_donor", "N_donor", "S_donor", "P_donor", "carboxyl", "amino", "thiol", "sulfonic", "phosphonic"]
PORE = ["BET_m2_g", "total_pore_volume_cm3_g", "average_pore_diameter_nm"]
ALGORITHMS = ["Ridge", "SVR", "RF", "ExtraTrees", "XGBoost", "CatBoost"]


def parse_args():
    parser = argparse.ArgumentParser(description="Train six no-T MOF regression models.")
    parser.add_argument("--input", required=True, help="Input .xlsx or .csv dataset")
    parser.add_argument("--output", required=True, help="Output directory")
    parser.add_argument("--target", required=True, help="Regression target column")
    parser.add_argument("--conditions", nargs="+", required=True, help="Condition columns to include")
    parser.add_argument("--target-label", default=None, help="Human-readable target label")
    parser.add_argument("--trials", type=int, default=30, help="Optuna trials per algorithm")
    return parser.parse_args()


def load_frame(path):
    return pd.read_csv(path) if path.suffix.lower() == ".csv" else pd.read_excel(path)


def add_multihot(df, column, prefix):
    tokens = sorted({t.strip() for value in df[column].fillna("none") for t in str(value).split(";") if t.strip()})
    names = []
    for token in tokens:
        name = f"{prefix}__{token}"
        df[name] = df[column].fillna("none").map(
            lambda value: int(token in {item.strip() for item in str(value).split(";") if item.strip()})
        )
        names.append(name)
    return names


def make_pipeline(numeric, estimator):
    preprocess = ColumnTransformer([
        ("main", Pipeline([
            ("impute", SimpleImputer(strategy="constant", fill_value="none")),
            ("onehot", OneHotEncoder(handle_unknown="ignore")),
        ]), ["main"]),
        ("numeric", SimpleImputer(strategy="median", add_indicator=True), numeric),
    ])
    return Pipeline([("preprocess", preprocess), ("model", estimator)])


def suggest_params(trial, name):
    if name == "Ridge":
        return {"alpha": trial.suggest_float("alpha", 0.01, 1000, log=True)}
    if name == "SVR":
        return {"C": trial.suggest_float("C", 0.1, 2000, log=True), "epsilon": trial.suggest_float("epsilon", 0.001, 10, log=True), "gamma": trial.suggest_float("gamma", 1e-4, 10, log=True)}
    if name in {"RF", "ExtraTrees"}:
        return {"n_estimators": trial.suggest_int("n_estimators", 300, 1000, step=100), "max_depth": trial.suggest_int("max_depth", 4, 30), "min_samples_leaf": trial.suggest_int("min_samples_leaf", 1, 8), "max_features": trial.suggest_float("max_features", 0.3, 1.0), "max_samples": trial.suggest_float("max_samples", 0.7, 1.0)}
    if name == "XGBoost":
        return {"n_estimators": trial.suggest_int("n_estimators", 300, 1000, step=100), "max_depth": trial.suggest_int("max_depth", 3, 10), "learning_rate": trial.suggest_float("learning_rate", 0.01, 0.2, log=True), "subsample": trial.suggest_float("subsample", 0.6, 1.0), "colsample_bytree": trial.suggest_float("colsample_bytree", 0.5, 1.0), "min_child_weight": trial.suggest_int("min_child_weight", 1, 10), "reg_alpha": trial.suggest_float("reg_alpha", 1e-4, 5, log=True), "reg_lambda": trial.suggest_float("reg_lambda", 0.1, 20, log=True)}
    return {"iterations": trial.suggest_int("iterations", 300, 1000, step=100), "depth": trial.suggest_int("depth", 4, 10), "learning_rate": trial.suggest_float("learning_rate", 0.01, 0.2, log=True), "l2_leaf_reg": trial.suggest_float("l2_leaf_reg", 0.1, 20, log=True), "random_strength": trial.suggest_float("random_strength", 0, 2), "bagging_temperature": trial.suggest_float("bagging_temperature", 0, 5)}


def make_estimator(name, params):
    if name == "Ridge": return Ridge(**params)
    if name == "SVR": return SVR(kernel="rbf", **params)
    if name == "RF": return RandomForestRegressor(random_state=SEED, n_jobs=-1, bootstrap=True, **params)
    if name == "ExtraTrees": return ExtraTreesRegressor(random_state=SEED, n_jobs=-1, bootstrap=True, **params)
    if name == "XGBoost": return XGBRegressor(objective="reg:squarederror", random_state=SEED, n_jobs=-1, **params)
    return CatBoostRegressor(loss_function="RMSE", random_seed=SEED, verbose=False, thread_count=-1, **params)


def group_name(encoded):
    if encoded.startswith("main__main_"): return "M"
    if encoded.startswith("numeric__metal__"): return "metal_node_elements"
    if encoded.startswith("numeric__linker__"): return "Linker family"
    if encoded.startswith("numeric__modulator__"): return "modulator"
    return encoded.replace("numeric__", "")


def component_name(encoded):
    if encoded.startswith("main__main_"): return encoded.removeprefix("main__main_")
    for prefix in ("numeric__metal__", "numeric__linker__", "numeric__modulator__"):
        if encoded.startswith(prefix): return encoded.removeprefix(prefix)
    return encoded.replace("numeric__", "")


def write_shap_book(summary, grouped, path):
    wb = Workbook(); first = wb.active; first.title = "Overall SHAP"
    for sheet, frame in [(first, summary)] + [(wb.create_sheet(name), data) for name, data in grouped.items()]:
        for row in [list(frame.columns)] + frame.values.tolist(): sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(name="Times New Roman", bold=True); cell.fill = PatternFill("solid", fgColor="D9EAF7"); cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(max(max(len(str(sheet.cell(row=row, column=column).value or "")) for row in range(1, sheet.max_row + 1)) + 2, 14), 34)
    wb.save(path)


def export_shap(pipe, x, folder, algorithm):
    transformed = pipe.named_steps["preprocess"].transform(x)
    encoded = list(pipe.named_steps["preprocess"].get_feature_names_out())
    model = pipe.named_steps["model"]
    if algorithm == "XGBoost":
        import xgboost
        values = model.get_booster().predict(xgboost.DMatrix(transformed), pred_contribs=True)[:, :-1]
    elif algorithm == "CatBoost":
        values = model.get_feature_importance(Pool(transformed), type="ShapValues")[:, :-1]
    else:
        values = shap.TreeExplainer(model).shap_values(transformed)
    frame = pd.DataFrame({"component": [component_name(v) for v in encoded], "feature_group": [group_name(v) for v in encoded], "mean_abs_shap": np.abs(np.asarray(values)).mean(axis=0)})
    summary = frame.groupby("feature_group", as_index=False)["mean_abs_shap"].sum().sort_values("mean_abs_shap", ascending=False).reset_index(drop=True)
    summary.insert(0, "rank", np.arange(1, len(summary) + 1)); summary["shap_percentage"] = 100 * summary["mean_abs_shap"] / summary["mean_abs_shap"].sum()
    grouped = {}
    for group, item in frame.groupby("feature_group"):
        if group in {"M", "metal_node_elements", "Linker family", "modulator"}:
            result = item[["component", "mean_abs_shap"]].sort_values("mean_abs_shap", ascending=False).reset_index(drop=True)
            result.insert(0, "rank_within_group", np.arange(1, len(result) + 1)); result["share_within_group_percent"] = 100 * result["mean_abs_shap"] / result["mean_abs_shap"].sum(); grouped[group] = result
    summary.to_csv(folder / "SHAP_分组占比.csv", index=False, encoding="utf-8-sig")
    write_shap_book(summary, grouped, folder / "SHAP_分组占比.xlsx")


def main():
    args = parse_args(); source = Path(args.input); output = Path(args.output); output.mkdir(parents=True, exist_ok=True)
    df = load_frame(source); df[args.target] = pd.to_numeric(df[args.target], errors="coerce"); df = df[df[args.target].notna()].copy().reset_index(drop=True)
    multi = add_multihot(df, "metal_node_elements", "metal") + add_multihot(df, "Linker family", "linker") + add_multihot(df, "modulator", "modulator")
    numeric = BINARY + multi + PORE + args.conditions; features = ["main"] + numeric
    x = df[features].copy()
    for column in numeric: x[column] = pd.to_numeric(x[column], errors="coerce").fillna(0)
    y = df[args.target].to_numpy(dtype=float); cv = KFold(n_splits=5, shuffle=True, random_state=CV_SEED); results = []
    for name in ALGORITHMS:
        folder = output / name; folder.mkdir(parents=True, exist_ok=True)
        objective = lambda trial: float(cross_val_score(make_pipeline(numeric, make_estimator(name, suggest_params(trial, name))), x, y, cv=cv, scoring="r2", n_jobs=1).mean())
        study = optuna.create_study(direction="maximize", sampler=optuna.samplers.TPESampler(seed=SEED)); study.optimize(objective, n_trials=args.trials, show_progress_bar=False)
        pipe = make_pipeline(numeric, make_estimator(name, study.best_params)); oof = cross_val_predict(pipe, x, y, cv=cv, method="predict", n_jobs=1); folds = []
        for fold, (train, test) in enumerate(cv.split(x), 1):
            pred = clone(pipe).fit(x.iloc[train], y[train]).predict(x.iloc[test]); folds.append({"fold": fold, "r2": float(r2_score(y[test], pred)), "mae": float(mean_absolute_error(y[test], pred)), "rmse": float(mean_squared_error(y[test], pred) ** 0.5), "n_validation": len(test)})
        metrics = {"model": name, "target": args.target, "target_label": args.target_label or args.target, "rows": len(df), "papers": int(df.paper_id.nunique()), "materials": int(df.material_id.nunique()), "r2_oof": float(r2_score(y, oof)), "mae_oof": float(mean_absolute_error(y, oof)), "rmse_oof": float(mean_squared_error(y, oof) ** 0.5), "r2_fold_mean": float(np.mean([v['r2'] for v in folds])), "r2_fold_sd": float(np.std([v['r2'] for v in folds], ddof=1)), "best_cv_r2": float(study.best_value), "validation": "ordinary random 5-fold KFold, shuffle=True, random_state=2026", "optimization": "Optuna TPE Bayesian optimization using the same 5-fold CV", "n_trials": args.trials, "T_feature": "removed"}
        pd.DataFrame({"paper_id": df.paper_id, "material_id": df.material_id, "observed": y, "oof_predicted": oof, "residual": y-oof}).to_csv(folder / "oof_predictions_5fold.csv", index=False, encoding="utf-8-sig")
        pd.DataFrame(folds).to_csv(folder / "fold_metrics.csv", index=False, encoding="utf-8-sig"); pd.DataFrame(study.trials_dataframe()).to_csv(folder / "bayesopt_history.csv", index=False, encoding="utf-8-sig")
        (folder / "metrics.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8"); (folder / "parameters.json").write_text(json.dumps({"all_input_features": features, "numeric_features": numeric, "target": args.target, "best_model_parameters": study.best_params}, ensure_ascii=False, indent=2), encoding="utf-8")
        pipe.fit(x, y); joblib.dump({"model": pipe, "feature_columns": features, "scaler": None}, folder / "final_pipeline.joblib")
        if name in {"RF", "ExtraTrees", "XGBoost", "CatBoost"}: export_shap(pipe, x, folder, name)
        results.append(metrics); print(f"{name}: OOF R2={metrics['r2_oof']:.4f}")
    pd.DataFrame(sorted(results, key=lambda item: item["r2_oof"], reverse=True)).to_csv(output / "model_comparison_bayesopt_cv5.csv", index=False, encoding="utf-8-sig")


if __name__ == "__main__":
    main()
